from io import BytesIO
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from docx import Document
from .drafting import ROOT, catalogue, draft
from .model import COMMERCIAL
from .schemas import DraftRequest, ExportRequest, Proposal

app = FastAPI(title="Proposal drafting assistant")
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")

@app.get("/", include_in_schema=False)
def home():
    return FileResponse(ROOT / "static/index.html")

@app.get("/api/catalogue")
def get_catalogue():
    return catalogue()

@app.post("/api/draft", response_model=Proposal)
def generate(req: DraftRequest):
    try:
        return draft(req)
    except ValueError as error:
        raise HTTPException(400, str(error)) from error

@app.post("/api/export")
def export(req: ExportRequest):
    proposal = generate(req.request)
    if any(COMMERCIAL.search(value) for value in [req.reviewer_notes, *req.edits.values()]):
        raise HTTPException(400, "Commercial edits must be made in an approved package pricing record, not reviewer notes.")
    for key, value in req.edits.items():
        proposal.sections[key] = ["Reviewer amendment — unconfirmed: " + value]
    doc = Document()
    doc.add_heading(proposal.title, 0)
    doc.add_paragraph("Client: " + (proposal.client_name or "[To be confirmed]"))
    doc.add_paragraph(proposal.fixed_wording)
    doc.add_paragraph(f"Template: {proposal.template_id} / version {proposal.template_version}. Approval: {proposal.template_approved}")
    doc.add_paragraph("Package: " + proposal.source_references[0]["id"] + " / version " + proposal.source_references[0]["version"] + " / " + proposal.source_references[0]["status"])
    for key in proposal.section_order:
        doc.add_heading(key.replace("_", " ").title(), level=1)
        for line in proposal.sections[key]:
            doc.add_paragraph(line, style="List Bullet")
    if req.reviewer_notes.strip():
        doc.add_heading("Reviewer amendments — unconfirmed", level=1)
        doc.add_paragraph(req.reviewer_notes)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": 'attachment; filename="proposal-draft.docx"'})
